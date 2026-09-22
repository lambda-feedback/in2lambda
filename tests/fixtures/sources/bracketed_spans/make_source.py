"""Writes the ``source.docx`` beside this file, which was run by hand.

python-docx is not a dependency of in2lambda and is not added as one: the document is
committed, and this script records what it holds so that anyone can write it again. The
import stands inside the main block because the test suite imports every module under
``tests`` to collect its doctests.

Word's underline, highlight and small capitals are the three formats pandoc's docx
reader turns into a bracketed span, which `in2lambda source add` unwraps.
"""

if __name__ == "__main__":
    from pathlib import Path

    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    document = Document()

    heading = document.add_heading("", level=1)
    heading.add_run("Hydraulic scale").underline = True

    paragraph = document.add_paragraph()
    label = paragraph.add_run("Question 2:")
    label.bold = True
    label.underline = True
    paragraph.add_run(" A hydraulic scale has two pistons joined by ")
    paragraph.add_run("oil").font.highlight_color = WD_COLOR_INDEX.YELLOW
    paragraph.add_run(".")

    note = document.add_paragraph()
    note.add_run("Note: ").font.small_caps = True
    note.add_run("the oil is incompressible.")

    # Beside this file, so that the fixture is written wherever the script is run from.
    document.save(Path(__file__).with_name("source.docx"))
